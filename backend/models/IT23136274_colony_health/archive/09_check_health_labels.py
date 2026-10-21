from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

BASE = Path(r"E:\MSPB_Dataset")

print("\n=== D2 HEALTH OBSERVATIONS ===")

wb = load_workbook(
    BASE / "D2_ant.xlsx",
    read_only=True,
    data_only=True
)

ws = wb["Sheet1"]
rows = list(ws.iter_rows(values_only=True))
headers = [str(x).strip() if x is not None else ""
           for x in rows[0]]

print("COLUMNS:", headers)

for column_name in [
    "Mortality cause",
    "status April 5",
    "status April 9",
    "status April 13",
    "status Apr 27",
    "status Apr 28"
]:
    if column_name not in headers:
        print("Column not found:", column_name)
        continue

    idx = headers.index(column_name)
    values = [
        str(row[idx]).strip()
        for row in rows[1:]
        if idx < len(row)
        and row[idx] is not None
        and str(row[idx]).strip()
    ]

    print("\n", column_name)
    print("Records with values:", len(values))

    for value, count in Counter(values).most_common():
        print(f"  {value}: {count}")

wb.close()

print("\n=== D1 PHENOTYPIC MEASUREMENTS ===")

wb = load_workbook(
    BASE / "D1_ant.xlsx",
    read_only=True,
    data_only=True
)

ws = wb["Phenotypic measurements"]

for row in ws.iter_rows(
    min_row=1,
    max_row=5,
    values_only=True
):
    print(row)

wb.close()

print("\nHealth label inspection completed.")
