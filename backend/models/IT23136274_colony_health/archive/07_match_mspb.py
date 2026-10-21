
import csv
from pathlib import Path
from openpyxl import load_workbook

BASE = Path(r"E:\MSPB_Dataset")

# Excel IDs read කිරීම
excel_ids = set()

for filename in ["D1_ant.xlsx", "D2_ant.xlsx"]:
    wb = load_workbook(BASE / filename, read_only=True, data_only=True)

    for sheet in wb.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())

        for col_index, header in enumerate(headers):
            if str(header).strip().lower() == "hive id":
                for row in rows:
                    if col_index < len(row) and row[col_index] is not None:
                        hive_id = str(row[col_index]).strip()
                        excel_ids.add(hive_id.zfill(5))
                break

    wb.close()

print("EXCEL HIVE IDs:", len(excel_ids))
print(sorted(excel_ids))

# Sensor IDs read කිරීම
for filename in ["D1_sensor_data.csv", "D2_sensor_data.csv"]:
    sensor_ids = set()

    with (BASE / filename).open(
        "r", encoding="utf-8-sig", newline=""
    ) as f:
        reader = csv.DictReader(f)

        for row in reader:
            tag = row["tag_number"].strip()
            if tag:
                sensor_ids.add(tag)

    # Candidate mapping only:
    # 202040 -> 02040
    candidate_ids = {
        tag[-5:] for tag in sensor_ids
        if len(tag) >= 5
    }

    matched = excel_ids.intersection(candidate_ids)

    print("\nFILE:", filename)
    print("Total sensor IDs:", len(sensor_ids))
    print("Candidate matched Hive IDs:", len(matched))
    print("Matches:", sorted(matched))
    print("Unmatched:", sorted(candidate_ids - excel_ids))

print("\nID comparison completed.")


print("\nEXCEL SHEET STRUCTURE")

for filename in ["D1_ant.xlsx", "D2_ant.xlsx"]:
    workbook = load_workbook(
        BASE / filename,
        read_only=True,
        data_only=True
    )

    print("\nFILE:", filename)

    for sheet in workbook.worksheets:
        print("\nSHEET:", sheet.title)
        print("ROWS:", sheet.max_row)

        for row in sheet.iter_rows(
            min_row=1,
            max_row=min(5, sheet.max_row),
            values_only=True
        ):
            print(row[:12])

    workbook.close()
