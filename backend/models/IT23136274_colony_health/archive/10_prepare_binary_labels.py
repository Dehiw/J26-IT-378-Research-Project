
import csv
from pathlib import Path
from collections import Counter
from datetime import datetime
from openpyxl import load_workbook

BASE = Path(r"E:\MSPB_Dataset")

LABELS = []
EXCLUDED = []

# 1. Beekeeper observations read කරන්න
wb = load_workbook(
    BASE / "D2_ant.xlsx",
    read_only=True,
    data_only=True
)

ws = wb["Sheet1"]
rows = ws.iter_rows(values_only=True)
headers = [str(x).strip() if x is not None else ""
           for x in next(rows)]

for row in rows:
    data = dict(zip(headers, row))

    hive_id = str(data.get("Hive ID") or "").strip()
    if not hive_id:
        continue

    # Most recent available observation
    observations = [
        ("2021-04-28", data.get("status Apr 28")),
        ("2021-04-27", data.get("status Apr 27"))
    ]

    selected = next(
        (
            (date, str(note).strip().lower())
            for date, note in observations
            if note is not None and str(note).strip()
        ),
        None
    )

    if selected is None:
        continue

    date, note = selected

    # Exclude colonies explicitly without beacon
    if "no beacon" in note:
        EXCLUDED.append((hive_id, note, "No beacon"))
        continue

    if any(word in note for word in [
        "queenless",
        "drone laying queen",
        "chalkbrood",
        "dead"
    ]):
        label = "Needs Inspection"
    elif note.startswith("good"):
        label = "Healthy"
    else:
        EXCLUDED.append((hive_id, note, "Unknown status"))
        continue

    # Candidate mapping only; NOT verified
    sensor_id = "2" + hive_id.zfill(5)

    LABELS.append({
        "hive_id": hive_id,
        "sensor_id": sensor_id,
        "date": date,
        "label": label,
        "note": note
    })

wb.close()

print("\n=== PROPOSED LABELS ===")
print("Total:", len(LABELS))
print("Classes:", dict(Counter(
    item["label"] for item in LABELS
)))
print("Excluded:", len(EXCLUDED))

# 2. Sensor readings වල date coverage බලන්න
wanted = {item["sensor_id"] for item in LABELS}
coverage = {}

with (BASE / "D2_sensor_data.csv").open(
    "r", encoding="utf-8-sig", newline=""
) as f:

    reader = csv.DictReader(f)

    for row in reader:
        tag = row["tag_number"].strip()

        if tag not in wanted:
            continue

        date = row["date"].strip()
        if not date:
            continue

        if tag not in coverage:
            coverage[tag] = {
                "first": date,
                "last": date,
                "count": 0
            }

        item = coverage[tag]
        item["first"] = min(item["first"], date)
        item["last"] = max(item["last"], date)
        item["count"] += 1

print("\n=== SENSOR COVERAGE ===")

for item in LABELS:
    sensor = coverage.get(item["sensor_id"])

    if sensor is None:
        print(item["hive_id"], item["label"],
              "NO SENSOR DATA")
        continue

    gap = (
        datetime.fromisoformat(item["date"])
        - datetime.fromisoformat(sensor["last"])
    ).days

    print(
        item["hive_id"],
        "|", item["label"],
        "|", "Last sensor:", sensor["last"],
        "|", "Gap:", gap, "days",
        "|", "Readings:", sensor["count"]
    )

print("\nSUMMARY")
print("Labeled colonies:", len(LABELS))
print("With candidate sensor match:",
      sum(x["sensor_id"] in coverage for x in LABELS))
print("Excluded observations:", len(EXCLUDED))
print("\nCompleted. No model trained.")


print("\n=== SENSOR GAP FILTER ===")

for max_gap in [7, 14, 21, 30]:
    usable = []

    for item in LABELS:
        sensor = coverage.get(item["sensor_id"])

        if sensor is None:
            continue

        gap = (
            datetime.fromisoformat(item["date"])
            - datetime.fromisoformat(sensor["last"])
        ).days

        if 0 <= gap <= max_gap:
            usable.append(item)

    counts = Counter(x["label"] for x in usable)

    print(f"\nMaximum gap: {max_gap} days")
    print("Usable candidates:", len(usable))
    print("Healthy:", counts.get("Healthy", 0))
    print("Needs Inspection:", counts.get("Needs Inspection", 0))

print("\nFiltering completed.")
