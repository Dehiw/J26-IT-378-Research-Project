
import csv
from pathlib import Path
from collections import Counter, defaultdict
from datetime import date

import joblib
import numpy as np
from openpyxl import load_workbook
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    precision_score,
    recall_score,
    f1_score,
    confusion_matrix,
    classification_report,
)

BASE = Path(__file__).parent
DATASET = Path(r"E:\MSPB_Dataset")
MODEL_PATH = BASE / "colony_sensor_anomaly_model_v2.joblib"

# Evaluation settings
START_DATE = date(2021, 4, 8)
END_DATE = date(2021, 4, 14)
MAX_GAP_DAYS = 21
MAX_SAMPLES_PER_HIVE = 500
MIN_READINGS = 25
MIN_DAYS = 2

# Experimental alert threshold, not scientifically validated.
ALERT_THRESHOLD = 0.05

rng = np.random.default_rng(42)

print("Loading trained Isolation Forest V2...")
model = joblib.load(MODEL_PATH)

# -------------------------------------------------
# 1. Read proposed inspection labels
# -------------------------------------------------

wb = load_workbook(
    DATASET / "D2_ant.xlsx",
    read_only=True,
    data_only=True,
)

ws = wb["Sheet1"]
rows = ws.iter_rows(values_only=True)

headers = [
    str(h).strip() if h is not None else ""
    for h in next(rows)
]

labels = {}

for row in rows:
    data = dict(zip(headers, row))

    hive_id = str(data.get("Hive ID") or "").strip()

    if not hive_id:
        continue

    observations = [
        (date(2021, 4, 28), data.get("status Apr 28")),
        (date(2021, 4, 27), data.get("status Apr 27")),
    ]

    selected = next(
        (
            (day, str(note).strip().lower())
            for day, note in observations
            if note is not None and str(note).strip()
        ),
        None,
    )

    if selected is None:
        continue

    inspection_date, note = selected

    if "no beacon" in note:
        continue

    if any(x in note for x in [
        "queenless",
        "drone laying queen",
        "chalkbrood",
        "dead",
    ]):
        label = 1  # Needs Inspection

    elif note.startswith("good"):
        label = 0  # Healthy candidate

    else:
        continue

    # CANDIDATE ID MAPPING: NOT YET VERIFIED!
    sensor_id = "2" + hive_id.zfill(5)

    labels[sensor_id] = {
        "hive_id": hive_id,
        "inspection_date": inspection_date,
        "label": label,
        "note": note,
    }

wb.close()

print("\nPROPOSED LABEL COUNTS")
print(dict(Counter(x["label"] for x in labels.values())))

# -------------------------------------------------
# 2. Read sensor data in a fixed historical window
# -------------------------------------------------

samples = defaultdict(list)
seen = Counter()
days_covered = defaultdict(set)
last_dates = {}

print("\nReading D2 sensor data...")

with (DATASET / "D2_sensor_data.csv").open(
    "r", encoding="utf-8-sig", newline=""
) as f:

    reader = csv.DictReader(f)

    for row in reader:
        sensor_id = row["tag_number"].strip()

        if sensor_id not in labels:
            continue

        try:
            reading_date = date.fromisoformat(
                row["date"].strip()
            )

            if not (START_DATE <= reading_date <= END_DATE):
                continue

            temp = float(row["temperature"])
            humidity = float(row["humidity"])

            if not (
                np.isfinite(temp)
                and np.isfinite(humidity)
                and -20 <= temp <= 60
                and 0 <= humidity <= 100
            ):
                continue

        except (ValueError, TypeError):
            continue

        seen[sensor_id] += 1
        days_covered[sensor_id].add(reading_date)

        if (
            sensor_id not in last_dates
            or reading_date > last_dates[sensor_id]
        ):
            last_dates[sensor_id] = reading_date

        value = [temp, humidity]
        reservoir = samples[sensor_id]

        # Random sampling with bounded memory
        if len(reservoir) < MAX_SAMPLES_PER_HIVE:
            reservoir.append(value)
        else:
            index = rng.integers(0, seen[sensor_id])

            if index < MAX_SAMPLES_PER_HIVE:
                reservoir[index] = value

# -------------------------------------------------
# 3. Predict one result per colony
# -------------------------------------------------

y_true = []
y_pred = []
results = []

print("\nCOLONY-LEVEL RESULTS")
print("-" * 85)

for sensor_id, info in sorted(labels.items()):

    if seen[sensor_id] < MIN_READINGS:
        continue

    if len(days_covered[sensor_id]) < MIN_DAYS:
        continue

    last_date = last_dates[sensor_id]

    gap_days = (
        info["inspection_date"] - last_date
    ).days

    if not (0 <= gap_days <= MAX_GAP_DAYS):
        continue

    X = np.asarray(samples[sensor_id], dtype=float)

    predictions = model.predict(X)

    anomaly_fraction = float(
        np.mean(predictions == -1)
    )

    # Fixed experimental colony-level decision
    predicted_label = int(
        anomaly_fraction > ALERT_THRESHOLD
    )

    actual_label = info["label"]

    y_true.append(actual_label)
    y_pred.append(predicted_label)

    actual_name = (
        "Needs Inspection"
        if actual_label == 1 else "Healthy"
    )
    predicted_name = (
        "Inspection Alert"
        if predicted_label == 1 else "Normal Pattern"
    )

    result = {
        "hive_id": info["hive_id"],
        "actual": actual_name,
        "predicted": predicted_name,
        "anomaly_percent": round(
            anomaly_fraction * 100, 2
        ),
        "gap_days": gap_days,
        "readings": seen[sensor_id],
    }

    results.append(result)

    print(
        f"{result['hive_id']} | "
        f"Actual: {actual_name} | "
        f"Predicted: {predicted_name} | "
        f"Anomaly: {result['anomaly_percent']}% | "
        f"Gap: {gap_days} days"
    )

# -------------------------------------------------
# 4. Evaluation metrics
# -------------------------------------------------

print("\n" + "=" * 55)
print("EXPLORATORY MSPB EVALUATION")
print("=" * 55)

if len(y_true) == 0:
    print("No usable matched colonies found.")
    raise SystemExit(0)

print("Evaluated colonies:", len(y_true))

print("Actual classes:", dict(Counter(y_true)))
print("Predicted classes:", dict(Counter(y_pred)))

print("\nAccuracy:",
      round(accuracy_score(y_true, y_pred) * 100, 2), "%")

print("Balanced Accuracy:",
      round(balanced_accuracy_score(y_true, y_pred) * 100, 2), "%")

print("Precision (Needs Inspection):",
      round(precision_score(
          y_true, y_pred, zero_division=0
      ), 3))

print("Recall (Needs Inspection):",
      round(recall_score(
          y_true, y_pred, zero_division=0
      ), 3))

print("F1-score (Needs Inspection):",
      round(f1_score(
          y_true, y_pred, zero_division=0
      ), 3))

print("\nCLASSIFICATION REPORT")
print(classification_report(
    y_true,
    y_pred,
    labels=[0, 1],
    target_names=["Healthy", "Needs Inspection"],
    zero_division=0,
))

print("\nCONFUSION MATRIX")
print("Rows: Actual | Columns: Predicted")
print("Order: Healthy, Needs Inspection")
print(confusion_matrix(
    y_true, y_pred, labels=[0, 1]
))

print("\nIMPORTANT LIMITATIONS")
print("- Candidate sensor-to-hive mapping is not verified.")
print("- Labels are based on beekeeper notes.")
print("- Sensor observations precede labels by 13+ days.")
print("- MSPB D2 was also used for model training.")
print("- This is NOT independent health accuracy.")
print("- Alert threshold is experimental.")
print("\nEvaluation completed. Model unchanged.")
